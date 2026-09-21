"""Descarga, normaliza y verifica los catálogos y clasificadores autónomos del INEGI que la API
expone como catálogos: SCIAN 2023/2018/2013, SINCO 2019/2011, CMO (histórica) y el Catálogo Único
de Claves de Áreas Geoestadísticas Estatales, Municipales y Localidades (AGEEML).

Origen de cada catálogo (URL exacta; todas verificadas con curl, sin navegador):
  SCIAN 2023  https://www.inegi.org.mx/contenidos/app/scian/scian_2023_categorias_y_productos.xlsx?v=20230719
              (es el "Estructura completa en:" que enlaza el visor https://www.inegi.org.mx/scian/, proyecto 14
              del API /app/api/clasificadores/interna_v2/estructura/info/proyecto/?proy=14). Cifras publicadas:
              Síntesis metodológica SCIAN 2023 (UPC 889463909682), cuadro "Total 20 94 305 610 1086".
  SCIAN 2018  https://www.inegi.org.mx/contenidos/app/scian/scian_2018_categorias_y_productos.xlsx
              Cifras publicadas: documento SCIAN 2018 (UPC 702825099695), cuadro comparativo 2018 vs 2013
              "Total 20 20 94 94 306 303 615 614 1084 1059".
  SCIAN 2013  https://www.inegi.org.mx/contenidos/app/clasificadores/scian/SCIAN_2013_estructura_productos.xlsx
              (proyecto 4 del mismo API). Cifras publicadas: mismo cuadro comparativo (20/94/303/614/1059).
  SINCO 2019  https://www.inegi.org.mx/contenidos/productos/prod_serv/contenidos/espanol/bvinegi/productos/nueva_estruc/702825198411.pdf
              (única forma publicada: PDF de 400 páginas, edición 2020, con fe de erratas 702825198411_2.pdf del
              2 de septiembre de 2025). No existe XLSX/CSV descargable. Se extrae la "Versión abreviada"
              (estructura) y la "Versión ampliada" (descripciones y ocupaciones). Cifras publicadas en el propio
              documento: 9 divisiones, 52 grupos principales, 163 subgrupos, 490 grupos unitarios.
  SINCO 2011  https://www.inegi.org.mx/contenidos/clasificadoresycatalogos/doc/sinco_tablas_comparativas.xlsx
              (hoja SINCO-CMO, "Actualización noviembre de 2012"; columnas de la estructura SINCO 2011). Cifras
              publicadas: el documento SINCO 2019 declara 468 grupos unitarios en 2011 y que 2019 tiene "un grupo
              principal menos, siete subgrupos y 22 grupos unitarios más" -> 9/53/156/468.
  CMO         https://www.inegi.org.mx/contenidos/clasificadoresycatalogos/doc/clasificacion_mexicana_de_ocupaciones_vol_i.pdf
              https://www.inegi.org.mx/contenidos/clasificadoresycatalogos/doc/clasificacion_mexicana_de_ocupaciones_vol_ii.pdf
              ("CMO - Histórica", volúmenes I y II; solo PDF). Estructura y títulos: listado "Versión Abreviada"
              del vol. II; descripciones: vol. I; ocupaciones individuales: sección 2 del vol. II. Cifras publicadas
              (vol. I, 1.4): 19 grupos principales y 461 grupos unitarios (no declara el número de subgrupos).
  AGEEML      https://www.inegi.org.mx/app/ageeml/  ->  "Catálogos completos":
                https://www.inegi.org.mx/contenidos/app/ageeml/catun_entidad.zip
                https://www.inegi.org.mx/contenidos/app/ageeml/catun_municipio.zip
                https://www.inegi.org.mx/contenidos/app/ageeml/catun_localidad.zip
              y el API del visor /app/api/cateml/catuniAPI/ (ObtenerFecCorLoc/Mun/Ent dan las fechas de corte y
              el total de registros por corte; EXPLOCAGEML exporta el catálogo de localidades vigente en CSV UTF-8).
              El zip nacional de localidades es del corte 2024/12 (296 814 localidades); el visor publica cortes
              mensuales, por eso las localidades normalizadas se toman del API en el corte más reciente y el zip
              queda como descarga masiva de referencia.

Salidas:
  data/catalogos/fuentes/<catalogo>/...        originales tal cual se descargaron (se reutilizan si ya existen)
  data/catalogos/normalizado/<catalogo>.csv    CSV UTF-8 con encabezado y columnas fijas (ver README.md)
  data/catalogos/verificacion.json             conteos obtenidos vs publicados por catálogo y nivel

Verificación: cada catálogo compara los conteos por nivel del CSV normalizado contra la cifra publicada por el
INEGI (diccionario PUBLICADO) y contra la integridad jerárquica (todo codigo_padre existe). El resultado se
imprime en tabla y se guarda en verificacion.json; una diferencia se reporta, no se corrige.

Uso: python3 scripts/catalogos_inegi.py [--solo scian_2023,sinco_2019,...] [--sin-descarga]
Requiere: openpyxl (con texto enriquecido), pymupdf; red para descargar la primera vez.
"""
import argparse, csv, io, json, pathlib, re, sys, time, unicodedata, zipfile
from collections import Counter, OrderedDict
from urllib.parse import urlencode
from urllib.request import Request, urlopen

RAIZ = pathlib.Path(__file__).resolve().parent.parent
BASE = RAIZ / 'data' / 'catalogos'; FUENTES = BASE / 'fuentes'; NORM = BASE / 'normalizado'
UA = 'Mozilla/5.0 (datosmexico.org; catalogos)'
INEGI = 'https://www.inegi.org.mx'

URLS = {
    'scian_2023': {'scian_2023_categorias_y_productos.xlsx': INEGI + '/contenidos/app/scian/scian_2023_categorias_y_productos.xlsx?v=20230719',
                   '889463909682_sintesis_metodologica_scian_2023.pdf': INEGI + '/contenidos/productos/prod_serv/contenidos/espanol/bvinegi/productos/nueva_estruc/889463909682.pdf'},
    'scian_2018': {'scian_2018_categorias_y_productos.xlsx': INEGI + '/contenidos/app/scian/scian_2018_categorias_y_productos.xlsx',
                   '702825099695_scian_2018.pdf': INEGI + '/contenidos/productos/prod_serv/contenidos/espanol/bvinegi/productos/nueva_estruc/702825099695.pdf'},
    'scian_2013': {'SCIAN_2013_estructura_productos.xlsx': INEGI + '/contenidos/app/clasificadores/scian/SCIAN_2013_estructura_productos.xlsx'},
    'sinco_2019': {'702825198411_sinco_2019.pdf': INEGI + '/contenidos/productos/prod_serv/contenidos/espanol/bvinegi/productos/nueva_estruc/702825198411.pdf',
                   '702825198411_2_fe_de_erratas.pdf': INEGI + '/contenidos/productos/prod_serv/contenidos/espanol/bvinegi/productos/nueva_estruc/fe_erratas/702825198411_2.pdf'},
    'sinco_2011': {'sinco_tablas_comparativas.xlsx': INEGI + '/contenidos/clasificadoresycatalogos/doc/sinco_tablas_comparativas.xlsx'},
    'cmo': {'clasificacion_mexicana_de_ocupaciones_vol_i.pdf': INEGI + '/contenidos/clasificadoresycatalogos/doc/clasificacion_mexicana_de_ocupaciones_vol_i.pdf',
            'clasificacion_mexicana_de_ocupaciones_vol_ii.pdf': INEGI + '/contenidos/clasificadoresycatalogos/doc/clasificacion_mexicana_de_ocupaciones_vol_ii.pdf'},
    'ageeml': {'catun_entidad.zip': INEGI + '/contenidos/app/ageeml/catun_entidad.zip',
               'catun_municipio.zip': INEGI + '/contenidos/app/ageeml/catun_municipio.zip',
               'catun_localidad.zip': INEGI + '/contenidos/app/ageeml/catun_localidad.zip'},
}
# Cifras publicadas por el INEGI (ver docstring para la fuente de cada una).
PUBLICADO = {
    'scian_2023': {'sector': 20, 'subsector': 94, 'rama': 305, 'subrama': 610, 'clase': 1086},
    'scian_2018': {'sector': 20, 'subsector': 94, 'rama': 306, 'subrama': 615, 'clase': 1084},
    'scian_2013': {'sector': 20, 'subsector': 94, 'rama': 303, 'subrama': 614, 'clase': 1059},
    'sinco_2019': {'division': 9, 'grupo_principal': 52, 'subgrupo': 163, 'grupo_unitario': 490},
    'sinco_2011': {'division': 9, 'grupo_principal': 53, 'subgrupo': 156, 'grupo_unitario': 468},
    'cmo': {'grupo_principal': 19, 'subgrupo': None, 'grupo_unitario': 461},
    'entidades': {'entidad': 32}, 'municipios': {'municipio': 2478},
}
VERIF = OrderedDict()

def log(*a): print(*a, flush=True)

def http(url, params=None, intentos=4, espera=5):
    """GET con reintentos; devuelve bytes."""
    if params: url = url + '?' + urlencode(params)
    for i in range(intentos):
        try:
            with urlopen(Request(url, headers={'User-Agent': UA}), timeout=600) as r: return r.read()
        except Exception as e:
            if i == intentos - 1: raise
            log(f'  reintento {i+1} ({e})'); time.sleep(espera * (i + 1))

def descargar(cat, sin_descarga=False):
    d = FUENTES / cat; d.mkdir(parents=True, exist_ok=True)
    for nombre, url in URLS[cat].items():
        p = d / nombre
        if p.exists() and p.stat().st_size > 0: continue
        if sin_descarga: raise SystemExit(f'falta {p} y se pidió --sin-descarga')
        log(f'  descargando {url}'); p.write_bytes(http(url)); log(f'  {p.name}: {p.stat().st_size:,} bytes')
    return d

def escribir_csv(nombre, cols, filas):
    NORM.mkdir(parents=True, exist_ok=True); p = NORM / nombre
    with open(p, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(cols)
        for r in filas: w.writerow([('' if v is None else v) for v in r])
    log(f'  -> {p.relative_to(RAIZ)} ({len(filas):,} filas)'); return p

def verificar(cat, conteos, publicado, extra=None):
    """Compara conteos obtenidos vs publicados; guarda y muestra el veredicto."""
    filas = []; ok = True
    for nivel, n in conteos.items():
        pub = (publicado or {}).get(nivel)
        estado = 'sin cifra publicada' if pub is None else ('cuadra' if pub == n else 'NO CUADRA')
        if estado == 'NO CUADRA': ok = False
        filas.append({'nivel': nivel, 'obtenido': n, 'publicado': pub, 'estado': estado})
        log(f'  {nivel:<16} obtenido {n:>7,}   publicado {str(pub) if pub is not None else "—":>7}   {estado}')
    VERIF[cat] = {'niveles': filas, 'cuadra': ok, **(extra or {})}
    return ok

def normalizar_txt(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)

# ----------------------------------------------------------------------------------------------- SCIAN
NIVELES_SCIAN = ['sector', 'subsector', 'rama', 'subrama', 'clase']
HOJAS_SCIAN = {'sector': 'SECTOR', 'subsector': 'SUBSECTOR', 'rama': 'RAMA', 'subrama': 'SUBRAMA', 'clase': 'CLASE'}

def celda_texto(v):
    """Texto plano de una celda y la marca en superíndice (T, CAN, EU) que el INEGI pone al final de los títulos."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    if v is None: return '', ''
    if isinstance(v, CellRichText):
        plano = ''; marca = ''
        for b in v:
            if isinstance(b, TextBlock):
                if b.font is not None and b.font.vertAlign == 'superscript': marca += b.text
                else: plano += b.text
            else: plano += str(b)
        return plano.strip(), marca.strip()
    return str(v).strip(), ''

def padre_scian(nivel, codigo, sectores):
    if nivel == 'sector': return ''
    if nivel == 'subsector':
        for s in sectores:
            partes = s.split('-')
            if len(partes) == 2 and partes[0] <= codigo[:2] <= partes[1]: return s
            if s == codigo[:2]: return s
        return codigo[:2]
    return codigo[:{'rama': 3, 'subrama': 4, 'clase': 5}[nivel]]

def normalizar_scian(cat, archivo):
    import openpyxl
    wb = openpyxl.load_workbook(archivo, read_only=True, rich_text=True)
    filas = []; productos = []; sectores = []
    for nivel in NIVELES_SCIAN:
        ws = wb[HOJAS_SCIAN[nivel]]; enc = None; ultimo = None
        for r in ws.iter_rows(values_only=False):
            vals = [c.value for c in r]
            if enc is None:
                if vals and celda_texto(vals[0])[0] == 'Código':
                    enc = [celda_texto(v)[0] for v in vals]
                    ix = {}
                    for i, h in enumerate(enc):
                        h = h.lower()
                        if h.startswith('código'): ix['codigo'] = i
                        elif h.startswith('título'): ix['titulo'] = i
                        elif h.startswith('descripción'): ix['descripcion'] = i
                        elif h.startswith('incluye'): ix['incluye'] = i
                        elif h.startswith('excluye'): ix['excluye'] = i
                        elif h.startswith('productos') or h.startswith('índice de bienes'): ix['productos'] = i
                continue
            cod = celda_texto(vals[ix['codigo']])[0] if ix['codigo'] < len(vals) else ''
            if cod:
                cod = cod.replace('.0', '') if re.fullmatch(r'\d+\.0', cod) else cod
                titulo, marca = celda_texto(vals[ix['titulo']])
                desc = celda_texto(vals[ix['descripcion']])[0] if 'descripcion' in ix else ''
                inc = celda_texto(vals[ix['incluye']])[0] if 'incluye' in ix and ix['incluye'] < len(vals) else ''
                exc = celda_texto(vals[ix['excluye']])[0] if 'excluye' in ix and ix['excluye'] < len(vals) else ''
                if nivel == 'sector': sectores.append(cod)
                filas.append([nivel, cod, padre_scian(nivel, cod, sectores), titulo, desc, inc, exc, marca]); ultimo = cod
            if nivel == 'clase' and 'productos' in ix and ix['productos'] < len(vals):
                prod = celda_texto(vals[ix['productos']])[0]
                if prod and ultimo:
                    prod = re.sub(r'^\s*-\s*', '', prod).strip()
                    if prod: productos.append([ultimo, prod])
    cols = ['nivel', 'codigo', 'codigo_padre', 'titulo', 'descripcion', 'incluye', 'excluye', 'marca']
    escribir_csv(f'{cat}.csv', cols, filas)
    escribir_csv(f'{cat}_productos.csv', ['codigo_clase', 'producto'], productos)
    codigos = {(f[0], f[1]) for f in filas}; codset = {f[1] for f in filas}
    huerfanos = [f[1] for f in filas if f[2] and f[2] not in codset]
    conteos = OrderedDict((n, sum(1 for f in filas if f[0] == n)) for n in NIVELES_SCIAN)
    verificar(cat, conteos, PUBLICADO[cat], {'productos': len(productos), 'huerfanos': huerfanos, 'duplicados': len(filas) - len(codigos)})
    if huerfanos: log(f'  AVISO: {len(huerfanos)} códigos con padre inexistente: {huerfanos[:10]}')

# ----------------------------------------------------------------------------------------------- PDF utilitario
def abrir_pdf(p):
    import pymupdf
    return pymupdf.open(str(p))

def lineas_pagina(page, columnas=False, margen=15, margen_x=0):
    """Líneas de texto de una página agrupando renglones a la misma altura (tolerancia de 4 pt). Con columnas=True
    separa las dos columnas de la página (izquierda primero) para reconstruir el flujo de lectura. Se descartan las
    líneas escritas en vertical (los pies de página girados de la CMO) y, con margen_x, las pegadas al borde."""
    mid = page.rect.width / 2; ancho = page.rect.width; porcol = {0: [], 1: []}
    for b in page.get_text('dict')['blocks']:
        if b.get('type') != 0: continue
        for l in b['lines']:
            if abs(l['dir'][0]) < 0.9: continue          # texto vertical
            x0, y0 = l['bbox'][0], l['bbox'][1]
            t = ''.join(sp['text'] for sp in l['spans']).strip()
            if not t: continue
            if margen_x and (x0 < margen_x or x0 > ancho - margen_x): continue
            col = (0 if x0 < mid - margen else 1) if columnas else 0
            porcol[col].append((y0, x0, t))
    out = []
    for col in (0, 1):
        ws = sorted(porcol[col]); linea = []; y_ref = None
        for y0, x0, t in ws:
            if y_ref is not None and y0 - y_ref > 4:
                linea.sort(); out.append((col, y_ref, linea[0][0], ' '.join(w[1] for w in linea))); linea = []
            if not linea: y_ref = y0
            linea.append((x0, t))
        if linea: linea.sort(); out.append((col, y_ref, linea[0][0], ' '.join(w[1] for w in linea)))
    return out

def unir_lineas(lineas):
    """Une renglones en un párrafo respetando guiones de corte de línea (incluido el guion suave)."""
    s = ''
    for l in lineas:
        l = l.strip()
        if not l: continue
        if s and (s.endswith('­') or (s.endswith('-') and l[:1].islower())): s = s[:-1] + l
        elif s: s += ' ' + l
        else: s = l
    return s.replace('­', '').replace('ﬁ', 'fi').replace('ﬂ', 'fl')

# ----------------------------------------------------------------------------------------------- SINCO 2019
NIV_SINCO = {1: 'division', 2: 'grupo_principal', 3: 'subgrupo', 4: 'grupo_unitario'}
# Fe de erratas 702825198411_2.pdf (INEGI, 2 de septiembre de 2025): ejemplos de ocupaciones del grupo unitario 5115.
ERRATAS_SINCO_2019 = {'5115': ['Preparador de bebidas', 'Barman', 'Bartender']}
PIE_SINCO = re.compile(r'^INEGI\. Sistema Nacional de Clasificaci')

def sinco2019_abreviada(doc):
    """Estructura ordenada (codigo, titulo) de la sección "Grupo unitario" de la Versión abreviada."""
    ini = fin = None; abrev = None
    for i, p in enumerate(doc):
        t = p.get_text('text')
        if abrev is None and 'Versión abreviada' in t: abrev = i; continue
        if abrev is not None and ini is None and re.search(r'^\d{4}\s+\S', t, re.M): ini = i   # primera página con grupos unitarios
        if ini is not None and 'Versión ampliada' in t: fin = i; break
    est = []
    for i in range(ini, fin):
        for l in doc[i].get_text('text').split('\n'):
            l = l.replace('\t', ' ').strip()
            if not l or PIE_SINCO.match(l) or l.isdigit() or l in ('División', 'Grupo principal', 'Subgrupo', 'Grupo unitario'): continue
            m = re.match(r'^(\d{1,4})\s*(\D.*)$', l)
            if m: est.append([m.group(1), m.group(2).strip()])
            elif est: est[-1][1] = unir_lineas([est[-1][1], l])
    return [(c, t.replace('ﬁ', 'fi').replace('ﬂ', 'fl')) for c, t in est]

def normalizar_sinco_2019(archivo):
    doc = abrir_pdf(archivo)
    abrev = sinco2019_abreviada(doc)
    # códigos en orden, sin repetidos (el PDF repite la clave 971 con dos títulos; se conserva la primera aparición
    # de cada clave y el título se toma de la Versión ampliada, que es la que describe cada categoría)
    orden = []; vistos = set(); tit_abrev = {}
    for c, t in abrev:
        if c not in vistos: vistos.add(c); orden.append(c); tit_abrev[c] = t
    # Versión ampliada: páginas entre "Versión ampliada" y "Anexo"
    ini = fin = None
    for i, p in enumerate(doc):
        t = p.get_text('text')
        if ini is None and 'Versión ampliada' in t: ini = i + 1
        if ini is not None and i > ini and re.search(r'Tabla de equivalencia SINCO 2011', t): fin = i; break
    lineas = []
    for i in range(ini, fin):
        for l in doc[i].get_text('text').split('\n'):
            l = l.replace('\t', ' ').strip()
            if not l or PIE_SINCO.match(l): continue
            lineas.append(l)   # se conservan los renglones solo numéricos: pueden ser una clave de encabezado (se descartan abajo)
    reg = OrderedDict(); pendientes = list(orden); actual = None; k = 0
    while k < len(lineas):
        l = lineas[k]; m = re.match(r'^(\d{1,4})\s+(\S.*)$', l)
        if not m and pendientes and l == pendientes[0] and k + 1 < len(lineas) and not lineas[k + 1][:1].isdigit() \
                and normalizar_txt(tit_abrev[l]).startswith(normalizar_txt(lineas[k + 1])[:12]):
            k += 1; m = re.match(r'^(\d{1,4})\s+(\S.*)$', f'{l} {lineas[k]}')   # clave sola y título en el renglón siguiente
        if m and pendientes and m.group(1) == pendientes[0]:
            cod = pendientes.pop(0); titulo = m.group(2).strip()
            objetivo = normalizar_txt(tit_abrev[cod])
            while normalizar_txt(titulo) != objetivo and objetivo.startswith(normalizar_txt(titulo)) and k + 1 < len(lineas) \
                    and not re.match(r'^(\d{1,4})\s+\S', lineas[k + 1]):
                k += 1; titulo = unir_lineas([titulo, lineas[k]])
            actual = reg[cod] = {'titulo': titulo, 'desc': [], 'ocup': [], 'seccion': 'desc'}
        elif l.isdigit(): pass   # número de página
        elif actual is not None:
            if re.fullmatch(r'(Ocupaciones|Funciones|Excluye|Incluye|Comprende)\s*:', l): actual['seccion'] = l.rstrip(':').lower()
            elif actual['seccion'] == 'desc': actual['desc'].append(l)
            elif actual['seccion'] == 'ocupaciones':
                if l.startswith('•'): actual['ocup'].append(l.lstrip('• ').strip())
                elif actual['ocup']: actual['ocup'][-1] = unir_lineas([actual['ocup'][-1], l])
        k += 1
    filas = []
    for cod in orden:
        r = reg.get(cod); nivel = NIV_SINCO[len(cod)]; padre = cod[:-1] if len(cod) > 1 else ''
        titulo = r['titulo'] if r else tit_abrev[cod]
        desc = unir_lineas(r['desc']) if r else ''
        ocup = [o.rstrip('.') for o in r['ocup']] if r else []
        nota = '' if r else 'sin ficha en la Versión ampliada'
        if cod in ERRATAS_SINCO_2019: ocup = ERRATAS_SINCO_2019[cod]; nota = 'ocupaciones según fe de erratas del 2 de septiembre de 2025'
        filas.append([nivel, cod, padre, titulo, desc, ' | '.join(ocup), nota])
    escribir_csv('sinco_2019.csv', ['nivel', 'codigo', 'codigo_padre', 'titulo', 'descripcion', 'ocupaciones', 'nota'], filas)
    codset = {f[1] for f in filas}; huerf = [f[1] for f in filas if f[2] and f[2] not in codset]
    conteos = OrderedDict((n, sum(1 for f in filas if f[0] == n)) for n in NIV_SINCO.values())
    dup = [c for c, n in Counter(c for c, _ in abrev).items() if n > 1]
    verificar('sinco_2019', conteos, PUBLICADO['sinco_2019'],
              {'huerfanos': huerf, 'sin_ficha_ampliada': [f[1] for f in filas if f[6]], 'sin_descripcion': sum(1 for f in filas if not f[4]),
               'claves_repetidas_en_abreviada': dup, 'no_encontrados_en_ampliada': pendientes})
    if huerf: log(f'  AVISO huérfanos: {huerf}')
    if dup: log(f'  AVISO claves repetidas en la Versión abreviada del PDF: {dup}')
    if pendientes: log(f'  AVISO códigos sin ficha en la Versión ampliada: {pendientes}')

# ----------------------------------------------------------------------------------------------- SINCO 2011
def normalizar_sinco_2011(archivo):
    import openpyxl
    ws = openpyxl.load_workbook(archivo, read_only=True)['SINCO-CMO']
    filas = []; vistos = set()
    for r in ws.iter_rows(min_row=3, values_only=True):
        c = r[3]
        if c is None: continue
        cod = str(c).strip()
        if not cod.isdigit() or cod in vistos: continue
        vistos.add(cod); t = str(r[4] or '').strip()
        filas.append([NIV_SINCO[len(cod)], cod, cod[:-1] if len(cod) > 1 else '', t, ''])
    escribir_csv('sinco_2011.csv', ['nivel', 'codigo', 'codigo_padre', 'titulo', 'descripcion'], filas)
    codset = {f[1] for f in filas}; huerf = [f[1] for f in filas if f[2] and f[2] not in codset]
    conteos = OrderedDict((n, sum(1 for f in filas if f[0] == n)) for n in NIV_SINCO.values())
    verificar('sinco_2011', conteos, PUBLICADO['sinco_2011'], {'huerfanos': huerf})

# ----------------------------------------------------------------------------------------------- CMO
NIV_CMO = {2: 'grupo_principal', 3: 'subgrupo', 4: 'grupo_unitario'}
PIE_CMO = re.compile(r'(Clasificaci[oó]n M[eé]xicana de Ocupaciones|INEGI\. Clasificación|Volumen (I|II)|^Grupo [Pp]rincipal \d\d$|^Versión Abreviada$|^Aspectos Generales)')
CABECERA_CMO = re.compile(r'\s*(Clasificaci[oó]n M[eé]xicana de Ocupaciones|Grupo [Pp]rincipal \d\d)\s*')

def limpiar_cmo(l):
    """Quita la cabecera corrida de página cuando quedó pegada a un renglón del cuerpo."""
    return CABECERA_CMO.sub(' ', l).strip() if CABECERA_CMO.search(l) and not PIE_CMO.fullmatch(l.strip()) else l.strip()

def es_mayusculas(s):
    letras = [c for c in s if c.isalpha()]
    return bool(letras) and all(c.isupper() for c in letras)

def cmo_estructura(doc2):
    """(codigo, titulo) de los tres niveles a partir del listado 'Versión Abreviada' del volumen II."""
    fin = None
    for i, p in enumerate(doc2):
        if i > 2 and re.search(r'2\. Clasificación Mexicana de\s*\nOcupaciones|GRUPO PRINCIPAL 11\s*\n\s*110\s*\n', p.get_text('text')): fin = i; break
    est = []
    for i in range(fin):
        for col, y, x, l in lineas_pagina(doc2[i], margen_x=25):
            l = limpiar_cmo(l)
            if not l or PIE_CMO.search(l) or l.isdigit(): continue
            m = re.match(r'^GRUPO PRINCIPAL (\d\d)$', l)
            if m: est.append([m.group(1), '']); continue
            m = re.match(r'^(\d{3,4})\s+(.+)$', l)
            if m: est.append([m.group(1), m.group(2).strip()]); continue
            if est and es_mayusculas(l): est[-1][1] = (est[-1][1] + ' ' + l).strip()
    return est

def cmo_descripciones(doc1):
    """Descripción de cada grupo principal / subgrupo / grupo unitario a partir del volumen I (dos columnas)."""
    ini = None
    for i, p in enumerate(doc1):
        if re.search(r'^\s*GRUPO PRINCIPAL 11\s*$', p.get_text('text'), re.M) and 'PROFESIONISTAS' in p.get_text('text'): ini = i; break
    flujo = []
    for i in range(ini, len(doc1)):
        for col, y, x, l in lineas_pagina(doc1[i], columnas=True, margen_x=25):
            l = limpiar_cmo(l)
            if not l or PIE_CMO.search(l) or l.isdigit(): continue
            flujo.append(l)
    desc = {}; actual = None; estado = None
    for l in flujo:
        m = re.match(r'^(GRUPO PRINCIPAL|SUBGRUPO|GRUPO UNITARIO)\s+(\d{2,4})$', l)
        if m: actual = desc[m.group(2)] = {'titulo': '', 'desc': []}; estado = 'titulo'; continue
        if actual is None: continue
        if estado == 'titulo':
            if es_mayusculas(l): actual['titulo'] = (actual['titulo'] + ' ' + l).strip(); continue
            estado = 'desc'
        if estado == 'desc':
            if re.match(r'^Algunas de sus (tareas|funciones)', l) or l.startswith('- '): estado = 'tareas'; continue
            actual['desc'].append(l)
    return {c: {'titulo': v['titulo'], 'desc': unir_lineas(v['desc'])} for c, v in desc.items()}

def cmo_ocupaciones(doc2):
    """Ocupaciones individuales por grupo unitario (sección 2 del volumen II, dos columnas)."""
    ini = None
    for i, p in enumerate(doc2):
        if i > 2 and re.search(r'2\. Clasificación Mexicana de\s*\nOcupaciones', p.get_text('text')): ini = i; break
    oc = []
    for i in range(ini, len(doc2)):
        for col, y, x, l in lineas_pagina(doc2[i], columnas=True, margen_x=25):
            l = limpiar_cmo(l)
            if not l or PIE_CMO.search(l) or l.isdigit(): continue
            m = re.match(r'^(\d{4})\s+(.+)$', l)
            if m:
                if es_mayusculas(m.group(2)): oc.append(None); continue   # título del grupo unitario, no una ocupación
                oc.append([m.group(1), m.group(2).strip()]); continue
            if re.match(r'^(GRUPO PRINCIPAL|\d{3})\b', l) or es_mayusculas(l): oc.append(None); continue
            if oc and oc[-1] is not None: oc[-1][1] = unir_lineas([oc[-1][1], l])
    return [o for o in oc if o]

def normalizar_cmo(d):
    doc1 = abrir_pdf(d / 'clasificacion_mexicana_de_ocupaciones_vol_i.pdf'); doc2 = abrir_pdf(d / 'clasificacion_mexicana_de_ocupaciones_vol_ii.pdf')
    est = cmo_estructura(doc2); desc = cmo_descripciones(doc1); ocup = cmo_ocupaciones(doc2)
    filas = []; vistos = set()
    for cod, tit in est:
        if cod in vistos: continue
        vistos.add(cod); dv = desc.get(cod, {})
        titulo = tit or dv.get('titulo', '')
        filas.append([NIV_CMO[len(cod)], cod, cod[:-1] if len(cod) > 2 else '', titulo, dv.get('desc', ''), dv.get('titulo', ''), '' if cod in desc else 'sin descripción en vol. I'])
    escribir_csv('cmo.csv', ['nivel', 'codigo', 'codigo_padre', 'titulo', 'descripcion', 'titulo_vol_i', 'nota'], filas)
    gu = {f[1] for f in filas if f[0] == 'grupo_unitario'}
    oc_ok = [o for o in ocup if o[0] in gu]; oc_sin = sorted({o[0] for o in ocup if o[0] not in gu})
    escribir_csv('cmo_ocupaciones.csv', ['codigo_grupo_unitario', 'ocupacion'], oc_ok)
    codset = {f[1] for f in filas}; huerf = [f[1] for f in filas if f[2] and f[2] not in codset]
    conteos = OrderedDict((n, sum(1 for f in filas if f[0] == n)) for n in NIV_CMO.values())
    solo_vol1 = sorted(set(desc) - codset)
    verificar('cmo', conteos, PUBLICADO['cmo'], {'huerfanos': huerf, 'ocupaciones': len(oc_ok), 'ocupaciones_con_clave_desconocida': oc_sin,
                                                'sin_descripcion': [f[1] for f in filas if f[6]], 'solo_en_vol_i': solo_vol1,
                                                'titulos_distintos_vol_i_vs_vol_ii': sum(1 for f in filas if f[5] and normalizar_txt(f[5]) != normalizar_txt(f[3]))})
    if huerf: log(f'  AVISO huérfanos: {huerf}')
    if solo_vol1: log(f'  AVISO claves con descripción en vol. I que no están en el listado del vol. II: {solo_vol1}')
    if oc_sin: log(f'  AVISO ocupaciones con clave fuera del catálogo: {oc_sin}')

# ----------------------------------------------------------------------------------------------- AGEEML
API_CAT = INEGI + '/app/api/cateml/catuniAPI/'

def csv_de_zip(zp):
    """Primer CSV del zip que decodifica como UTF-8 (los zips traen ANSI y UTF-8 con el mismo contenido)."""
    with zipfile.ZipFile(zp) as z:
        nombres = [n for n in z.namelist() if n.lower().endswith('.csv')]
        for pref in ('utf8', 'utf', ''):
            for n in nombres:
                if pref in n.lower():
                    raw = z.read(n)
                    try: txt = raw.decode('utf-8'); enc = 'utf-8'
                    except UnicodeDecodeError: txt = raw.decode('latin-1'); enc = 'latin-1'
                    if pref or enc == 'utf-8': return n, enc, txt
        n = nombres[0]; return n, 'latin-1', z.read(n).decode('latin-1')

def leer_csv_ageeml(txt):
    """Lee el CSV del AGEEML. Las coordenadas sexagesimales traen comillas sin escapar (21°52´47.362"N), así que se
    separa por el patrón '","' en vez de usar el módulo csv."""
    lineas = [l for l in txt.replace('\r\n', '\n').split('\n') if l.strip()]
    enc = [h.strip() for h in lineas[0].split(',')]
    filas = []
    for l in lineas[1:]:
        l = l.strip()
        if l.startswith('"') and l.endswith('"'): l = l[1:-1]
        filas.append(l.split('","'))
    return enc, filas

def fecha_iso(ddmmyyyy):
    d, m, y = ddmmyyyy[:10].split('/'); return f'{y}-{m}-{d}'

def normalizar_ageeml(d, sin_descarga=False):
    api = d / 'api'; api.mkdir(exist_ok=True)
    # fechas de corte y totales publicados por el visor
    cortes = {}
    for k, ep in (('ent', 'ObtenerFecCorEnt'), ('mun', 'ObtenerFecCorMun'), ('loc', 'ObtenerFecCorLoc')):
        p = api / f'fechas_corte_{k}.json'
        if not p.exists():
            if sin_descarga: raise SystemExit(f'falta {p}')
            p.write_bytes(http(API_CAT + ep + '/'))
        cortes[k] = [(fecha_iso(x['nombre']), int(x['valor'])) for x in json.loads(p.read_text(encoding='utf-8'))]
    ult_loc, tot_loc = cortes['loc'][0]; ult_mun, tot_mun = cortes['mun'][0]; ult_ent, tot_ent = cortes['ent'][0]
    log(f'  cortes vigentes según el API: entidades {ult_ent} ({tot_ent}), municipios {ult_mun} ({tot_mun}), localidades {ult_loc} ({tot_loc:,})')

    # entidades (zip)
    n, enc, txt = csv_de_zip(d / 'catun_entidad.zip'); h, filas = leer_csv_ageeml(txt); ix = {c: i for i, c in enumerate(h)}
    ent = [[r[ix['CVE_ENT']], r[ix['NOM_ENT']], r[ix['NOM_ABR']], r[ix['POB_TOTAL']], r[ix['POB_MASCULINA']], r[ix['POB_FEMENINA']], r[ix['TOTAL DE VIVIENDAS HABITADAS']], ult_ent] for r in filas]
    escribir_csv('entidades.csv', ['cve_ent', 'nom_ent', 'nom_abr', 'pob_total', 'pob_masculina', 'pob_femenina', 'viviendas_habitadas', 'fecha_corte'], ent)
    verificar('entidades', OrderedDict(entidad=len(ent)), {'entidad': tot_ent}, {'archivo_zip': n, 'fecha_corte': ult_ent})

    # municipios (zip)
    n, enc, txt = csv_de_zip(d / 'catun_municipio.zip'); h, filas = leer_csv_ageeml(txt); ix = {c: i for i, c in enumerate(h)}
    mun = [[r[ix['CVE_ENT']], r[ix['CVE_MUN']], r[ix['CVEGEO']], r[ix['NOM_ENT']], r[ix['NOM_ABR']], r[ix['NOM_MUN']], r[ix['CVE_CAB']], r[ix['NOM_CAB']],
            r[ix['POB_TOTAL']], r[ix['POB_MASCULINA']], r[ix['POB_FEMENINA']], r[ix['TOTAL DE VIVIENDAS HABITADAS']], ult_mun] for r in filas]
    escribir_csv('municipios.csv', ['cve_ent', 'cve_mun', 'cvegeo', 'nom_ent', 'nom_abr', 'nom_mun', 'cve_cab', 'nom_cab', 'pob_total', 'pob_masculina', 'pob_femenina', 'viviendas_habitadas', 'fecha_corte'], mun)
    verificar('municipios', OrderedDict(municipio=len(mun), entidad=len({m[0] for m in mun})), {'municipio': tot_mun, 'entidad': 32}, {'archivo_zip': n, 'fecha_corte': ult_mun})

    cols_loc = ['cve_ent', 'cve_mun', 'cve_loc', 'cvegeo', 'nom_ent', 'nom_abr', 'nom_mun', 'nom_loc', 'ambito', 'lat', 'lon', 'lat_dms', 'lon_dms', 'altitud', 'cve_carta',
                'pob_total', 'pob_masculina', 'pob_femenina', 'viviendas_habitadas', 'fecha_corte']
    def fila_loc(r, ix, corte):
        return [r[ix['CVE_ENT']], r[ix['CVE_MUN']], r[ix['CVE_LOC']], r[ix['CVEGEO']], r[ix['NOM_ENT']], r[ix['NOM_ABR']], r[ix['NOM_MUN']], r[ix['NOM_LOC']], r[ix['AMBITO']],
                r[ix['LAT_DECIMAL']], r[ix['LON_DECIMAL']], r[ix['LATITUD']], r[ix['LONGITUD']], r[ix['ALTITUD']], r[ix['CVE_CARTA']],
                r[ix['POB_TOTAL']], r[ix['POB_MASCULINA']], r[ix['POB_FEMENINA']], r[ix['TOTAL DE VIVIENDAS HABITADAS']], corte]

    # localidades: descarga masiva (zip nacional) -> corte deducido por el total contra la serie del API
    n, enc, txt = csv_de_zip(d / 'catun_localidad.zip'); h, filas = leer_csv_ageeml(txt); ix = {c: i for i, c in enumerate(h)}
    cortes_zip = [f for f, v in cortes['loc'] if v == len(filas)]
    corte_zip = max(cortes_zip) if cortes_zip else ''
    locz = [fila_loc(r, ix, corte_zip) for r in filas]
    escribir_csv('localidades_descarga_masiva.csv', cols_loc, locz)
    verificar('localidades_descarga_masiva', OrderedDict(localidad=len(locz), municipio=len({(l[0], l[1]) for l in locz}), entidad=len({l[0] for l in locz})),
              {'localidad': len(filas) if cortes_zip else None, 'municipio': tot_mun, 'entidad': 32},
              {'archivo_zip': n, 'codificacion': enc, 'cortes_del_api_con_ese_total': cortes_zip, 'fecha_corte_asignada': corte_zip,
               'ambito': dict(Counter(l[8] for l in locz)), 'cvegeo_unicos': len({l[3] for l in locz})})

    # localidades vigentes en el corte más reciente, exportadas por el API en una sola consulta nacional (CSV UTF-8,
    # nombres en minúscula con acento = forNom 3, sin bajas = union 0). El API devuelve el catálogo nacional
    # completo aunque se pida una entidad, por eso se consulta una sola vez con entidad=00.
    p = api / f'localidades_00_{ult_loc}.csv'
    if not p.exists():
        if sin_descarga: raise SystemExit(f'falta {p}')
        log(f'  API localidades nacional corte {ult_loc}')
        params = {'nquery': 1, 'entidad': '00', 'busqueda': 'cve_ent=00', 'fecha': f'{ult_loc[8:10]}/{ult_loc[5:7]}/{ult_loc[:4]}', 'forNom': 3,
                  'ri': 0, 'rf': tot_loc, 'orden': '0 asc', 'union': 0, 'type_': 2, 'encodingType': 1}
        p.write_bytes(http(API_CAT + 'EXPLOCAGEML/', params))
    h, filas = leer_csv_ageeml(p.read_text(encoding='utf-8-sig')); ix = {c: i for i, c in enumerate(h)}
    loc = [fila_loc(r, ix, ult_loc) for r in filas]
    escribir_csv('localidades.csv', cols_loc, loc)
    verificar('localidades', OrderedDict(localidad=len(loc), municipio=len({(l[0], l[1]) for l in loc}), entidad=len({l[0] for l in loc})),
              {'localidad': tot_loc, 'municipio': tot_mun, 'entidad': 32},
              {'fecha_corte': ult_loc, 'ambito': dict(Counter(l[8] for l in loc)), 'cvegeo_unicos': len({l[3] for l in loc}),
               'municipios_del_catalogo_sin_localidad': sorted({m[2] for m in mun} - {l[0] + l[1] for l in loc})})

# ----------------------------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--solo', default=''); ap.add_argument('--sin-descarga', action='store_true')
    a = ap.parse_args(); solo = {s.strip() for s in a.solo.split(',') if s.strip()}
    pasos = [('scian_2023', lambda d: normalizar_scian('scian_2023', d / 'scian_2023_categorias_y_productos.xlsx')),
             ('scian_2018', lambda d: normalizar_scian('scian_2018', d / 'scian_2018_categorias_y_productos.xlsx')),
             ('scian_2013', lambda d: normalizar_scian('scian_2013', d / 'SCIAN_2013_estructura_productos.xlsx')),
             ('sinco_2019', lambda d: normalizar_sinco_2019(d / '702825198411_sinco_2019.pdf')),
             ('sinco_2011', lambda d: normalizar_sinco_2011(d / 'sinco_tablas_comparativas.xlsx')),
             ('cmo', normalizar_cmo),
             ('ageeml', lambda d: normalizar_ageeml(d, a.sin_descarga))]
    for cat, fn in pasos:
        if solo and cat not in solo: continue
        log(f'== {cat}'); d = descargar(cat, a.sin_descarga); fn(d)
    BASE.mkdir(parents=True, exist_ok=True)
    (BASE / 'verificacion.json').write_text(json.dumps(VERIF, ensure_ascii=False, indent=1), encoding='utf-8')
    log('\nResumen:')
    for cat, v in VERIF.items():
        print(f"  {cat:<28} {'cuadra' if v['cuadra'] else 'NO CUADRA':<10} " + '  '.join(f"{n['nivel']}={n['obtenido']}/{n['publicado'] if n['publicado'] is not None else '—'}" for n in v['niveles']))

if __name__ == '__main__': main()
